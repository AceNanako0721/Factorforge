from pathlib import Path
import re, json, hashlib, shutil
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
DOC = ROOT / 'doc'
STAGE = Path('runtime/doc-build/v11-build/srs')
STAGE.mkdir(parents=True, exist_ok=True)
ASSETS = DOC / 'assets'
ASSETS.mkdir(exist_ok=True)
FONT = 'C:/Windows/Fonts/msyh.ttc'
FONT_B = 'C:/Windows/Fonts/msyhbd.ttc'

def set_font(style,size,bold=False,font='Microsoft YaHei'):
    style.font.name=font; style.font.size=Pt(size); style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr()
    rf=rp.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rp.insert(0,rf)
    for key in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn('w:'+key),font)

def field(p,instr):
    r=p.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); r._r.append(begin)
    r=p.add_run(); node=OxmlElement('w:instrText'); node.set(qn('xml:space'),'preserve'); node.text=instr; r._r.append(node)
    r=p.add_run(); sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); r._r.append(sep)
    p.add_run('')
    r=p.add_run(); end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); r._r.append(end)

def inline(p,text):
    chunks=re.split(r'(\[[^\]]+\]\(https?://[^)]+\))',text)
    for chunk in chunks:
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',chunk)
        if m:
            h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r'); prop=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'254D6D'); prop.append(color); r.append(prop)
            t=OxmlElement('w:t'); t.text=m[1]; r.append(t); h.append(r); p._p.append(h)
        else:
            # Requirement identifiers provide visual scanning anchors.
            m=re.match(r'((?:REQ-[A-Z]+-\d+|NFR-\d+)〔[^〕]+〕)(.*)',chunk)
            if m: p.add_run(m[1]).bold=True; p.add_run(m[2])
            else: p.add_run(chunk)

source=(DOC/'要求分析式样书_v1.1.md').read_text(encoding='utf-8')
# A rerun removes generated contents and retains the canonical marker.
source=re.sub(r'<!-- TOC -->.*?<!-- /TOC -->','<!-- TOC -->',source,flags=re.S)
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(1.9)
sec.left_margin=Cm(2.0); sec.right_margin=Cm(2.0)
sec.header_distance=Cm(.75); sec.footer_distance=Cm(.85)
normal=doc.styles['Normal']; set_font(normal,10.5)
normal.paragraph_format.line_spacing=1.16
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.widow_control=True
for sty,size in [('Title',23),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
    st=doc.styles[sty]; set_font(st,size,True)
    st.paragraph_format.space_before=Pt(13 if sty!='Title' else 34)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True
for i in [1,2,3]:
    name=f'TOC {i}'
    if name not in doc.styles: doc.styles.add_style(name,1)
    st=doc.styles[name]; set_font(st,10 if i==1 else 9.5,i==1)
    st.paragraph_format.space_before=Pt(3 if i==1 else 0)
    st.paragraph_format.space_after=Pt(1)
    st.paragraph_format.line_spacing=1.0
    st.paragraph_format.left_indent=Cm((i-1)*.55)
for name in ['Header','Footer']:
    set_font(doc.styles[name],8)
    doc.styles[name].paragraph_format.space_after=Pt(0)
sec.different_first_page_header_footer=True
header=sec.header.paragraphs[0]; header.text=''
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 '); field(footer,' PAGE '); footer.add_run(' 页')
settings=doc.settings.element
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for node in settings.findall(qn('w:evenAndOddHeaders')):
    settings.remove(node)
upd=OxmlElement('w:updateFields'); upd.set(qn('w:val'),'true'); settings.append(upd)
doc.core_properties.title='要求分析式样书_v1.1'
doc.core_properties.subject='需求分析与验证基线'
doc.core_properties.author='Factorforge'
doc.core_properties.keywords='FF-SRS-001,需求分析,情绪池,风险控制,负反馈'

expected=[]; images=[]; headings=[]
lines=source.splitlines(); i=0; cover=True
while i<len(lines):
    line=lines[i].strip(); i+=1
    if not line: continue
    if line=='<!-- TOC -->':
        p=doc.add_paragraph(); field(p,' TOC \\o "1-2" \\h \\z \\u ')
        doc.add_page_break(); continue
    if line.startswith('# '):
        text=line[2:]; p=doc.add_paragraph(style='Title')
        subject,ending=text.split('要求分析式样书')
        p.add_run(subject).add_break(); p.add_run('要求分析式样书'+ending)
        expected.append(text); continue
    if line=='## 目录':
        doc.add_page_break(); p=doc.add_paragraph('目录'); set_font(doc.styles['TOC Heading'],16,True)
        p.style=doc.styles['TOC Heading']
        cover=False; continue
    m=re.match(r'^(#{2,4}) (.*)$',line)
    if m:
        level=len(m[1])-1; text=m[2]
        p=doc.add_paragraph(text,f'Heading {level}'); headings.append((level,text)); expected.append(text)
        continue
    m=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line)
    if m:
        p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(6)
        r=p.add_run(); shape=r.add_picture(str(DOC/m[2]),width=Cm(16.4))
        shape._inline.docPr.set('descr',m[1]); images.append(m[2])
        p=doc.add_paragraph(m[1]); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(6)
        for r in p.runs: r.font.size=Pt(9)
        expected.append(m[1]); continue
    if line.startswith('|'):
        block=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i].strip()); i+=1
        rows=[]
        for line2 in block:
            vals=[v.strip() for v in line2.strip('|').split('|')]
            if all(re.fullmatch(r':?-+:?',v) for v in vals): continue
            rows.append(vals)
        cols=len(rows[0]); table=doc.add_table(rows=0,cols=cols)
        table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
        if cols==2: widths=[4.0,13.0]
        elif cols==4: widths=[2.0,4.9,4.5,5.6]
        else: widths=[3.3,5.5,8.2]
        # Dedicated compact schemas.
        if rows[0][0]=='参数族': widths=[3.1,8.1,5.8]
        elif rows[0][0]=='用例': widths=[3.1,5.1,8.8]
        elif rows[0][0]=='用户需求章节': widths=[4.0,4.5,8.5]
        elif rows[0][0]=='编号': widths=[2.1,7.8,3.5,3.6]
        elif rows[0][0]=='版本': widths=[1.2,3.2,9.6,3.0]
        for col,w in zip(table.columns,widths): col.width=Cm(w)
        props=table._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'D9D9D9'); borders.append(edge)
        props.append(borders)
        margins=OxmlElement('w:tblCellMar')
        for side,v in [('top',90),('bottom',90),('left',100),('right',100)]:
            el=OxmlElement('w:'+side); el.set(qn('w:w'),str(v)); el.set(qn('w:type'),'dxa'); margins.append(el)
        props.append(margins)
        for ri,vals in enumerate(rows):
            row=table.add_row(); trpr=row._tr.get_or_add_trPr()
            nosplit=OxmlElement('w:cantSplit'); trpr.append(nosplit)
            if ri==0:
                repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
            for ci,(cell,text) in enumerate(zip(row.cells,vals)):
                expected.append(re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)',r'\1',text)); cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EDF2' if ri==0 else ('F7F9FA' if ri%2==0 else 'FFFFFF')); cell._tc.get_or_add_tcPr().append(shade)
                p=cell.paragraphs[0]; inline(p,text)
                if (rows[0][0]=='工作包 / 主责' and ci==1) or (rows[0][0] in ['准入门','编号','变更号'] and ci==0): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1.1
                if ri==0: p.paragraph_format.keep_with_next=True
                for r in p.runs: r.font.size=Pt(9); r.bold=(ri==0)
        spacer=doc.add_paragraph(); spacer.paragraph_format.space_after=Pt(4); spacer.paragraph_format.space_before=Pt(0); spacer.paragraph_format.line_spacing=Pt(4)
        continue
    p=doc.add_paragraph(); inline(p,line)
    if cover:
        p.paragraph_format.space_after=Pt(13)
    expected.append(re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)',r'\1',line))

out=STAGE/'requirements.docx'; doc.save(out)
(STAGE/'expected.json').write_text(json.dumps(expected,ensure_ascii=False,indent=2),encoding='utf-8')
(STAGE/'headings.json').write_text(json.dumps(headings,ensure_ascii=False,indent=2),encoding='utf-8')
(STAGE/'canonical.md').write_text(source,encoding='utf-8')
(STAGE/'source-sha256.txt').write_text(hashlib.sha256((DOC/'用户需求.md').read_bytes()).hexdigest(),encoding='ascii')
print(json.dumps({'docx':str(out),'paragraph_cell_units':len(expected),'headings':len(headings),'images':len(images)},ensure_ascii=False))
