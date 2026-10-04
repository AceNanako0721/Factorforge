"""Build seven Word documents from canonical Markdown, preserving body order."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib, json, re, shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[2]
V=ROOT/'doc/v2.0'
STAGE=Path('runtime/doc-build/factorforge-v20')
STAGE.mkdir(parents=True,exist_ok=True)
FILES=[p for p in sorted(V.glob('*.md')) if p.name not in ('README.md','v1.1需求迁移与待决事项.md')]

def font(style,size,bold=False):
    style.font.name='Microsoft YaHei'; style.font.size=Pt(size); style.font.bold=bold
    style.font.color.rgb=RGBColor.from_string('1D2939')
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ('ascii','hAnsi','eastAsia'): fonts.set(qn('w:'+attr),'Microsoft YaHei')

def plain(text): return text.replace('**','').replace('`','')

def inline(paragraph,text):
    cursor=0
    for m in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)',text):
        paragraph.add_run(plain(text[cursor:m.start()]))
        label,target=m.groups()
        if target.endswith('.md') and (V/target).is_file(): target=target[:-3]+'.docx'
        rel=paragraph.part.relate_to(target,RT.HYPERLINK,is_external=True)
        link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
        run=OxmlElement('w:r');rp=OxmlElement('w:rPr')
        color=OxmlElement('w:color');color.set(qn('w:val'),'176B91');rp.append(color)
        run.append(rp);t=OxmlElement('w:t');t.text=plain(label);run.append(t);link.append(run)
        paragraph._p.append(link);cursor=m.end()
    paragraph.add_run(plain(text[cursor:]))
    return plain(re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',text))

def shade(cell,color):
    tcpr=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd');s.set(qn('w:fill'),color);tcpr.append(s)

manifest=[]
for n,path in enumerate(FILES,1):
    d=Document();s=d.sections[0]
    # Runtime's default style can contain a title underline; remove it.
    for border in d.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.7);s.left_margin=Cm(2);s.right_margin=Cm(2)
    s.footer_distance=Cm(.7)
    font(d.styles['Normal'],10.5)
    normal=d.styles['Normal'].paragraph_format
    normal.space_after=Pt(5);normal.line_spacing=1.12
    font(d.styles['Title'],21,True)
    for level,size in [(1,15),(2,12),(3,11)]:
        style=d.styles[f'Heading {level}'];font(style,size,True)
        style.paragraph_format.space_before=Pt(10 if level==1 else 7)
        style.paragraph_format.space_after=Pt(4)
        style.paragraph_format.keep_with_next=True
    footer=s.footer.paragraphs[0];footer.alignment=2
    r=footer.add_run('Factorforge v2.0  ·  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(90,100,115)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    d.core_properties.title=path.stem;d.core_properties.author='Factorforge';d.core_properties.last_modified_by='Factorforge'
    d.core_properties.created=datetime(2026,10,4,tzinfo=timezone.utc)
    d.core_properties.modified=datetime(2026,10,4,tzinfo=timezone.utc)
    lines=path.read_text(encoding='utf-8').splitlines();expected=[];i=0
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.startswith('| '):
            rows=[]
            while i<len(lines) and lines[i].startswith('| '):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x) for x in row):rows.append(row)
                i+=1
            cols=len(rows[0]);assert all(len(row)==cols for row in rows)
            table=d.add_table(rows=0,cols=cols);table.autofit=False
            widths=([3.0,8.2,5.8] if cols==3 else [5.4,11.6])
            if cols==3 and rows[0][0]=='式样': widths=[2.1,10,4.9]
            if cols==3 and rows[0][0]=='编号':widths=[2.1,6,8.9]
            if cols==3 and rows[0][0]=='阶段':widths=[3.0,8,6]
            for col,w in zip(table.columns,widths):col.width=Cm(w)
            for rownum,values in enumerate(rows):
                row=table.add_row()
                trpr=row._tr.get_or_add_trPr()
                nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
                if rownum==0:trpr.append(OxmlElement('w:tblHeader'))
                for cell,value,width in zip(row.cells,values,widths):
                    cell.width=Cm(width)
                    p=cell.paragraphs[0];expected.append(inline(p,value))
                    p.paragraph_format.keep_with_next=(rownum==0)
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.line_spacing=1.06
                    for run in p.runs:
                        run.font.size=Pt(9)
                        if rownum==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
                    shade(cell,'233E58' if rownum==0 else ('F0F4F8' if rownum%2 else 'FFFFFF'))
            d.add_paragraph().paragraph_format.space_after=Pt(2)
            continue
        m=re.match(r'^(#{1,4}) (.+)',line)
        if m:
            level=len(m[1]);p=d.add_paragraph(style='Title' if level==1 else f'Heading {level-1}')
            expected.append(inline(p,m[2]));i+=1;continue
        chunk=[line];i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','| ')):
            chunk.append(lines[i]);i+=1
        p=d.add_paragraph();expected.append(inline(p,' '.join(chunk)))
    slug=f'doc-{n:02}'
    local=STAGE/(slug+'.docx');d.save(local)
    (STAGE/(slug+'.expected.json')).write_text(json.dumps(expected,ensure_ascii=False),encoding='utf-8')
    manifest.append({'slug':slug,'md':str(path),'md_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'local_docx':str(local),'output_docx':str(path.with_suffix('.docx')),'expected_units':len(expected)})
(STAGE/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Built {len(manifest)} canonical Word documents')
